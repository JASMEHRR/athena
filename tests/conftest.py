"""Shared fixtures: an isolated Athena environment and synthetic decks."""

from __future__ import annotations

import json
import shutil
from pathlib import Path

import pymupdf
import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from athena import config, render

REAL_CONTENT = config.ROOT / "content"


@pytest.fixture
def env(tmp_path, monkeypatch):
    """Point every Athena path at a temp folder. Returns a namespace of paths."""
    content = tmp_path / "content"
    data = tmp_path / "data"
    reports = tmp_path / "reports"
    college = tmp_path / "college"
    inbox = tmp_path / "inbox"
    for d in (content, data, reports, college, inbox):
        d.mkdir()
    subjects = json.loads((REAL_CONTENT / "subjects.json").read_text(encoding="utf-8"))
    for subject in subjects:
        subject["deck_ids"] = []  # the test environment starts with no decks
    (content / "subjects.json").write_text(json.dumps(subjects), encoding="utf-8")
    (content / "sources.json").write_text(json.dumps({"skip_dirs": [], "files": []}), encoding="utf-8")
    paths = {
        "CONTENT_DIR": content,
        "DATA_DIR": data,
        "REPORTS_DIR": reports,
        "COLLEGE_DIR": college,
        "INBOX_DIR": inbox,
        "SLIDES_DIR": data / "slides",
        "CONVERTED_DIR": data / "converted",
        "REVIEW_QUEUE_DIR": data / "review_queue",
        "BACKUP_DIR": data / "backups",
        "CACHE_DIR": data / "cache",
        "DB_PATH": data / "athena.db",
    }
    for name, value in paths.items():
        monkeypatch.setattr(config, name, value)
    # Never launch PowerPoint from tests.
    monkeypatch.setattr(render, "_powerpoint_available", lambda: False)
    monkeypatch.setattr(config, "soffice_path", lambda: None)

    class Env:
        pass

    e = Env()
    for name, value in paths.items():
        setattr(e, name.lower(), value)
    e.root = tmp_path
    return e


def make_pptx(path: Path, extra_bullet: str = "") -> Path:
    """A small deck: title, bullets with levels, table, chart, picture, end."""
    path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Module 9 Pricing"
    s.placeholders[1].text = "Dr. Test Professor"

    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Price Elasticity"
    body = s.placeholders[1].text_frame
    body.text = "Elasticity measures how demand responds to price"
    p = body.add_paragraph()
    p.text = "Elastic demand: value above 1, as at Maruti in 2019"
    p.level = 1
    p = body.add_paragraph()
    p.text = "Inelastic demand: value below 1 for salt and medicines in daily use by households"
    p.level = 1
    if extra_bullet:
        p = body.add_paragraph()
        p.text = extra_bullet
    s.notes_slide.notes_text_frame.text = "Ask the class about petrol prices."

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Elasticity table"
    rows = [["Good", "Elasticity"], ["Salt", "0.1"], ["Cars", "1.8"]]
    table = s.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Sales by quarter"
    data = CategoryChartData()
    data.categories = ["Q1", "Q2"]
    data.add_series("Sales", (10.0, 12.5))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(4), data)

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Demand curve"
    png = path.parent / "pic.png"
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 40), False)
    pix.set_rect(pix.irect, (200, 30, 30))
    pix.save(str(png))
    s.shapes.add_picture(str(png), Inches(2), Inches(1.5), Inches(8), Inches(5))

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Thank you"

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))
    return path


def make_pdf(path: Path) -> Path:
    """A 3-page slide-shaped PDF: title page, bullet page, sparse page."""
    doc = pymupdf.open()
    page = doc.new_page(width=960, height=540)
    page.insert_text((80, 250), "Consumer Behaviour", fontsize=40)
    page.insert_text((80, 310), "Dr. Test Professor", fontsize=20)

    page = doc.new_page(width=960, height=540)
    page.insert_text((60, 70), "Buying Decision Process", fontsize=32)
    y = 140
    for indent, text in [
        (0, "• Need recognition starts the process for every buyer"),
        (0, "• Information search follows when the need is strong"),
        (30, "– Personal sources such as family and friends"),
        (30, "– Commercial sources such as advertising and dealers"),
        (0, "• Evaluation of alternatives and the final purchase decision"),
    ]:
        page.insert_text((80 + indent, y), text, fontsize=18)
        y += 40

    page = doc.new_page(width=960, height=540)
    page.insert_text((60, 70), "Model", fontsize=32)
    page.draw_rect(pymupdf.Rect(100, 120, 400, 400), color=(0, 0, 1))

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    doc.close()
    return path
