"""Extract a .pptx deck into slides: titles, bullets with levels, tables, notes,
chart data, SmartArt text and pictures.

Pictures are saved under data/slides/<deck_id>/ (deduplicated by content hash).
Page renders are added separately by athena.render.
"""

from __future__ import annotations

import hashlib
import logging
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Emu

from . import config
from .extract_common import assemble_slide, clean_line
from .models import Bullet, Chart, ChartSeries, Slide

log = logging.getLogger(__name__)

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}

# A picture covering at least this share of the slide is treated as content.
BIG_PICTURE_SHARE = 0.10
# This many plain drawn shapes (boxes, arrows) on one slide suggests a diagram.
DIAGRAM_SHAPES = 6


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def _shape_type(shape):
    """python-pptx raises for shapes it does not recognise; treat those as unknown."""
    try:
        return shape.shape_type
    except (NotImplementedError, AttributeError, KeyError):
        return None


def _paragraph_text(p_element) -> str:
    """Text of one a:p, including equation text (m:t) and line breaks."""
    parts: list[str] = []
    for el in p_element.iter():
        name = _local(el.tag)
        if name == "t" and el.text:
            parts.append(el.text)
        elif name == "br":
            parts.append(" ")
    return clean_line("".join(parts))


def _text_frame_bullets(text_frame) -> list[Bullet]:
    bullets: list[Bullet] = []
    for paragraph in text_frame.paragraphs:
        text = _paragraph_text(paragraph._p)
        if text:
            bullets.append(Bullet(level=min(int(paragraph.level or 0), 8), text=text))
    return bullets


def _table_rows(table) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            texts = [_paragraph_text(p._p) for p in cell.text_frame.paragraphs]
            cells.append(" ".join(t for t in texts if t))
        if any(cells):
            rows.append(cells)
    return rows


def _chart_data(chart) -> Chart | None:
    try:
        plot = chart.plots[0]
        categories = [str(c) for c in plot.categories]
        series = [
            ChartSeries(name=str(s.name or ""), values=[None if v is None else float(v) for v in s.values])
            for s in plot.series
        ]
        title = ""
        if chart.has_title and chart.chart_title.has_text_frame:
            title = clean_line(chart.chart_title.text_frame.text)
        return Chart(title=title, chart_type=str(chart.chart_type), categories=categories, series=series)
    except Exception as exc:  # python-pptx raises varied errors on unusual charts
        log.debug("chart unreadable: %s", exc)
        return None


def _smartart_text(shape, slide_part) -> list[str]:
    """SmartArt keeps its words in a separate diagram-data part."""
    lines: list[str] = []
    for rel_ids in shape._element.iter(f"{{{DIAGRAM_URI}}}relIds"):
        rid = rel_ids.get(f"{{{NS_R}}}dm")
        if not rid:
            continue
        try:
            part = slide_part.related_part(rid)
            root = etree.fromstring(part.blob)
        except Exception as exc:
            log.debug("SmartArt data unreadable: %s", exc)
            continue
        for p in root.iter(f"{{{NS_A}}}p"):
            text = _paragraph_text(p)
            if text:
                lines.append(text)
    return lines


class _SlideReader:
    def __init__(self, slide, slide_area: int, image_dir: Path, image_hashes: dict[str, str]):
        self.slide = slide
        self.slide_area = max(slide_area, 1)
        self.image_dir = image_dir
        self.image_hashes = image_hashes
        self.title = ""
        self.bullets: list[Bullet] = []
        self.tables: list[list[list[str]]] = []
        self.charts: list[Chart] = []
        self.images: list[str] = []
        self.picture_heavy = False
        self.diagram = False
        self.plain_shapes = 0

    def read(self) -> None:
        title_shape = None
        try:
            title_shape = self.slide.shapes.title
        except (AttributeError, KeyError):
            title_shape = None
        if title_shape is not None and title_shape.has_text_frame:
            self.title = " ".join(
                t for t in (_paragraph_text(p._p) for p in title_shape.text_frame.paragraphs) if t
            )
        self._walk(self.slide.shapes, skip=title_shape)
        if not self.title:
            self._guess_title()
        if self.plain_shapes >= DIAGRAM_SHAPES:
            self.diagram = True

    def _guess_title(self) -> None:
        """No title placeholder: use a short first line as the title."""
        if self.bullets and len(self.bullets[0].text.split()) <= 12:
            self.title = self.bullets.pop(0).text

    def _walk(self, shapes, skip=None) -> None:
        ordered = sorted(shapes, key=lambda s: ((s.top or 0) // Emu(300000), s.left or 0))
        for shape in ordered:
            if skip is not None and shape.shape_id == skip.shape_id:
                continue
            try:
                self._read_shape(shape)
            except Exception as exc:  # one odd shape must not lose the whole slide
                log.debug("shape %s unreadable: %s", getattr(shape, "name", "?"), exc)

    def _read_shape(self, shape) -> None:
        shape_type = _shape_type(shape)
        if shape_type == MSO_SHAPE_TYPE.GROUP:
            self._walk(shape.shapes)
            return
        if getattr(shape, "has_table", False) and shape.has_table:
            rows = _table_rows(shape.table)
            if rows:
                self.tables.append(rows)
            return
        if getattr(shape, "has_chart", False) and shape.has_chart:
            chart = _chart_data(shape.chart)
            if chart is not None:
                self.charts.append(chart)
            return
        graphic = shape._element.find(f".//{{{NS_A}}}graphicData")
        if graphic is not None and graphic.get("uri") == DIAGRAM_URI:
            self.diagram = True
            for line in _smartart_text(shape, self.slide.part):
                self.bullets.append(Bullet(level=0, text=line))
            return
        if shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
            self._save_picture(shape)
            return
        if getattr(shape, "has_text_frame", False):
            found = _text_frame_bullets(shape.text_frame)
            if found:
                if shape.is_placeholder and shape.placeholder_format.type in TITLE_TYPES and not self.title:
                    self.title = " ".join(b.text for b in found)
                else:
                    self.bullets.extend(found)
                return
        if shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.LINE, None):
            self.plain_shapes += 1

    def _save_picture(self, shape) -> None:
        try:
            image = shape.image
            blob = image.blob
        except Exception as exc:
            log.debug("picture unreadable: %s", exc)
            return
        area = (shape.width or 0) * (shape.height or 0)
        if area / self.slide_area >= BIG_PICTURE_SHARE:
            self.picture_heavy = True
        else:
            return  # logos and icons are not worth keeping
        digest = hashlib.sha256(blob).hexdigest()[:16]
        if digest not in self.image_hashes:
            ext = (image.ext or "png").lower()
            target = self.image_dir / f"img-{digest}.{ext}"
            target.parent.mkdir(parents=True, exist_ok=True)
            if not target.exists():
                target.write_bytes(blob)
            self.image_hashes[digest] = target.relative_to(config.DATA_DIR).as_posix()
        rel = self.image_hashes[digest]
        if rel not in self.images:
            self.images.append(rel)


def _notes(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    frame = slide.notes_slide.notes_text_frame
    if frame is None:
        return ""
    lines = [_paragraph_text(p._p) for p in frame.paragraphs]
    return "\n".join(line for line in lines if line)


def extract_pptx(path: Path, deck_id: str) -> tuple[list[Slide], str]:
    """Return (slides, deck title). Raises on files python-pptx cannot open."""
    prs = Presentation(str(path))
    slide_area = int(prs.slide_width or 0) * int(prs.slide_height or 0)
    image_dir = config.SLIDES_DIR / deck_id
    image_hashes: dict[str, str] = {}
    raw = list(prs.slides)
    slides: list[Slide] = []
    for index, slide in enumerate(raw, start=1):
        reader = _SlideReader(slide, slide_area, image_dir, image_hashes)
        reader.read()
        slides.append(
            assemble_slide(
                n=index,
                total=len(raw),
                title=reader.title,
                bullets=reader.bullets,
                tables=reader.tables,
                notes=_notes(slide),
                charts=reader.charts,
                images=reader.images,
                picture_heavy=reader.picture_heavy,
                diagram=reader.diagram,
            )
        )
    title = slides[0].title if slides and slides[0].title else path.stem
    return slides, title
